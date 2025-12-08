import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font,Alignment

f_class = pd.read_csv(r'./source_data/FunctionalClass.csv')
sign_system = pd.read_csv(r'./source_data/SignSys.csv')

county_list = f_class['County'].unique().tolist()

writer =pd.ExcelWriter('./output/2025_VMT_Report.xlsx')


#conditions for f_system classifications
interstate_urban = (f_class['FedFSystemID'] ==1) & (f_class['Urbanization'] == 'Urban')
freeways_and_expressways_urban = (f_class['FedFSystemID'] ==2) & (f_class['Urbanization'] == 'Urban')
other_principal_arterial_urban = (f_class['FedFSystemID'] ==3) & (f_class['Urbanization'] == 'Urban')
minor_arterial_urban = (f_class['FedFSystemID'] ==4) & (f_class['Urbanization'] == 'Urban')
major_collector_urban = (f_class['FedFSystemID'] ==5) & (f_class['Urbanization'] == 'Urban')
minor_collector_urban = (f_class['FedFSystemID'] ==6) & (f_class['Urbanization'] == 'Urban')
local_urban = (f_class['FedFSystemID'] ==7) & (f_class['Urbanization'] == 'Urban')
interstate_rural = (f_class['FedFSystemID'] ==1) & (f_class['Urbanization'] == 'Rural')
freeways_and_expressways_rural = (f_class['FedFSystemID'] ==2) & (f_class['Urbanization'] == 'Rural')
other_principal_arterial_rural = (f_class['FedFSystemID'] ==3) & (f_class['Urbanization'] == 'Rural')
minor_arterial_rural = (f_class['FedFSystemID'] ==4) & (f_class['Urbanization'] == 'Rural')
major_collector_rural = (f_class['FedFSystemID'] ==5) & (f_class['Urbanization'] == 'Rural')
minor_collector_rural = (f_class['FedFSystemID'] ==6) & (f_class['Urbanization'] == 'Rural')
local_rural = (f_class['FedFSystemID'] ==7) & (f_class['Urbanization'] == 'Rural')
total = (f_class['FedFSystemID'] ==99) & (f_class['Urbanization'] == 'All')

#conditions for sign system
interstate = (sign_system['SignSys_Code'] == 1)
united_states = (sign_system['SignSys_Code'] == 2)
west_virginia = (sign_system['SignSys_Code'] == 3)
county_route = (sign_system['SignSys_Code'] == 4)
state_park = (sign_system['SignSys_Code'] == 6)
fans = (sign_system['SignSys_Code'] == 7)
harp = (sign_system['SignSys_Code'] == 8)
total_sign = (sign_system['SignSys_Code'] == 99)



choices = ['Interstate Urban','Freeways and Expressways Urban','Other Principal Arterial Urban','Minor Arterial Urban','Major Collector Urban','Minor Collector Urban','Local Urban',
           'Interstate Rural','Freeways and Expressways Rural','Other Principal Arterial Rural','Minor Arterial Rural','Major Collector Rural','Minor Collector Rural','Local Rural','Total - Functional Classification']

choices_sign = ['Interstate','US','WV','CO','State Park & Forest','FANS','HARP','Total - Sign System']


conditions = [interstate_urban,
freeways_and_expressways_urban,
other_principal_arterial_urban,
minor_arterial_urban,
major_collector_urban,
minor_collector_urban,
local_urban,interstate_rural,
freeways_and_expressways_rural,
other_principal_arterial_rural,
minor_arterial_rural,
major_collector_rural,
minor_collector_rural,
local_rural,total]

conditions_sign = [interstate,
united_states,
west_virginia,
county_route,
state_park,
fans,
harp,
total_sign
]

f_class['functional_class'] = np.select(conditions,choices,default = 'Failed')
sign_system['sign_system'] = np.select(conditions_sign,choices_sign,default='Failed')

# print(f_class['functional_class'].unique())

f_class_beta = f_class[['functional_class','County','Miles','Miles_Percentage','Annual_Vehicle_Miles_Millions',
                 'Daily_Vehicle_Miles_Travel_Thousands','VMT_Percentage','Lane_Miles','Lanes_Percentage']]
sign_system_beta = sign_system[['sign_system','County','Miles','Miles_Percentage','Annual_Vehicle_Miles_Millions',
                 'Daily_Vehicle_Miles_Travel_Thousands','VMT_Percentage','Lane_Miles','Lanes_Percentage']]

# print(f_class_beta)
# print(sign_system_beta)


# mat1 = np.vstack([f_class_beta.columns,f_class_beta.to_numpy()])
# mat2 = np.vstack([sign_system_beta.columns,sign_system_beta.to_numpy()])


for county in county_list:
   f_class_fil = f_class_beta[f_class_beta['County'] == county]
   sign_system_fil = sign_system_beta[sign_system_beta['County']==county]
   mat1 = np.vstack([f_class_fil.columns,f_class_fil.to_numpy()])
   mat2 = np.vstack([sign_system_fil.columns,sign_system_fil.to_numpy()])
   full_report = pd.concat([pd.DataFrame(mat2),pd.DataFrame(mat1)]).reset_index()
   full_report.columns = full_report.iloc[0]
   full_report.drop(['County'],inplace=True,axis=1)
   full_report.drop([0],inplace=True,axis=1)
   full_report.drop(full_report.index[0],inplace=True,axis=0)

   full_report.loc[9,'sign_system'] = 'Functional Classification'
   full_report.loc[9,['Miles','Miles_Percentage','Annual_Vehicle_Miles_Millions',
                 'Daily_Vehicle_Miles_Travel_Thousands','VMT_Percentage','Lane_Miles','Lanes_Percentage']] = ''
   full_report.to_excel(writer,sheet_name = f'{county}',index = False)

writer.close()


wb = openpyxl.load_workbook('./output/2025_VMT_Report.xlsx')
for ws in wb.worksheets:
   #inserting and styling Title of Sheet
   ws.insert_rows(idx=1,amount=4)
   ws.cell(row=1,column=1).value = 'Traffic Analysis 1 Report - 2025'
   ws.cell(row=1,column=1).font = Font(bold=True,underline="single",size = 16)
   ws.merge_cells(start_row=1, start_column=1,end_row=1,end_column=8)
   ws.cell(row=1,column=1).alignment = Alignment(horizontal='center')
   #styling and adding county
   ws.cell(row=2,column=1).value =ws.title
   ws.cell(row=2,column=1).font = Font(bold=True,underline="single",size = 16)
   ws.merge_cells(start_row=2, start_column=1,end_row=2,end_column=8)
   ws.cell(row=2,column=1).alignment = Alignment(horizontal='center')
   #styling sign_system
   ws.cell(row=4,column=1).value = 'Sign System'
   ws.cell(row=5,column=1).value = ''
   ws.cell(row=4,column=1).font = Font(bold=True,underline="single")
   ws.merge_cells(start_row=4, start_column=1,end_row=4,end_column=8)
   ws.cell(row=4,column=1).alignment = Alignment(horizontal='center')
   #styling Functional Class
   ws.cell(row=14,column=1).font = Font(bold=True,underline="single")
   ws.merge_cells(start_row=14, start_column=1,end_row=14,end_column=8)
   ws.cell(row=14,column=1).alignment = Alignment(horizontal='center')
   #changing column names
   ws.cell(row=5,column=2).value = 'Miles'
   ws.cell(row=5,column=3).value = '% of County Total Miles'
   ws.cell(row=5,column=4).value = 'Annual Vehicle Miles (Millions)'
   ws.cell(row=5,column=5).value = 'Daily Vehicle Miles (Thousands)'
   ws.cell(row=5,column=6).value = '% of County Total VMT'
   ws.cell(row=5,column=7).value = 'Lane Miles'
   ws.cell(row=5,column=8).value = '% of Lane Total Miles'
   # for col in ws.columns:
   #   max_length = 0
   #   column = col[0].column_letter # Get the column name
   #   for cell in col:
   #       try: # Necessary to avoid error on empty cells
   #           if len(str(cell.value)) > max_length:
   #               max_length = len(str(cell.value))
   #       except:
   #           pass
   #   adjusted_width = (max_length + 2) * 1.2
   #   ws.column_dimensions[column].width = adjusted_width
wb.save('./output/2025_VMT_Report.xlsx')